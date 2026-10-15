#!/usr/bin/env python3
"""Import faux / PHI-cleaned cases for case-based-learning experiments (docs/CASE_IMPORT.md).

Input: the template workbook (.xlsx) or a directory / list of per-sheet CSVs (Cases_Medical.csv, Cases_NonMedical.csv,
Prior_Responses.csv, Outcomes.csv, Field_Reliability.csv).

  1. validate: schema, required columns, allowed values, cross-sheet links, and a PHI guard over every cell.
     Any error -> nothing is written, exit 2.
  2. convert: per-domain app items (items/<domain>.json), an app-loadable encounters.csv for the AI tele-doc domain
     (Domains/AIDocReview/EncounterImporter.swift, `-importCSV <path>` or Documents/encounters.csv), and unified
     decision records (sdr.decision.v1, provenance=imported_unsigned, never signed / never in the device ledger).
  3. --split: a learning experiment. Train = cases with day_offset < --cutoff (+ their prior HITM responses and outcomes);
     test = cases with day_offset >= cutoff, blinded (answers held out). Two conditions over the SAME test cases:
       no_history   : the test case alone
       with_history : the test case + train cases/responses/outcomes (same subject first)
     manifest.json records sha256 of every file so the pack is frozen.

Usage: python3 tools/import_cases.py <xlsx|dir|csv...> [--out DIR] [--split] [--cutoff 30] [--validate-only] [--export]
"""
from __future__ import annotations
import argparse, csv, datetime as dt, hashlib, json, os, re, sys
sys.path.insert(0, os.path.dirname(__file__))
from case_schema import SHEETS, OPTIONS_BY_DOMAIN, RELIABILITY_APP, APP_FIELDS  # noqa: E402

CASE_SHEETS = ("Cases_Medical", "Cases_NonMedical")
FORBIDDEN_COLS = re.compile(r"(^|_)(name|first|last|surname|mrn|ssn|dob|birth|address|street|city|zip|postal|phone|fax|email|"
                            r"insurance|member_id|account|license|plate|vin|device_serial|serial|ip|url|photo|date)($|_)", re.I)
MONTHS = r"(jan(uary)?|feb(ruary)?|mar(ch)?|apr(il)?|may|jun(e)?|jul(y)?|aug(ust)?|sep(t(ember)?)?|oct(ober)?|nov(ember)?|dec(ember)?)"
PHI = [
 ("email", re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")),
 ("url", re.compile(r"https?://|www\.", re.I)),
 ("ip address", re.compile(r"\b\d{1,3}(\.\d{1,3}){3}\b")),
 ("ssn", re.compile(r"\b\d{3}-\d{2}-\d{4}\b")),
 ("phone", re.compile(r"(\+?1[\s.-]?)?\(?\b\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}\b")),
 ("date", re.compile(r"\b\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}\b|\b(19|20)\d{2}[/.-]\d{1,2}[/.-]\d{1,2}\b")),
 ("date", re.compile(rf"\b{MONTHS}\.?\s+\d{{1,2}}(st|nd|rd|th)?\b|\b\d{{1,2}}(st|nd|rd|th)?\s+(of\s+)?{MONTHS}\b", re.I)),
 ("dob / birth", re.compile(r"\b(dob|d\.o\.b|born|birth\s*date|date of birth)\b", re.I)),
 ("mrn / record number", re.compile(r"\b(mrn|medical record|record\s*(no|number|#)|acct|account\s*(no|number|#)|member\s*id|policy\s*(no|#))\b", re.I)),
 ("long id number", re.compile(r"\b\d{7,}\b")),
 ("street address", re.compile(r"\b\d{1,6}\s+([A-Z][a-z]+\s+){1,3}(st|street|ave|avenue|rd|road|blvd|lane|ln|dr|drive|ct|court|way|pl|place|hwy)\b\.?", re.I)),
 ("zip code", re.compile(r"\b\d{5}-\d{4}\b|\b(zip|postal)\s*(code)?\s*:?\s*\d{3,5}\b", re.I)),
 ("name", re.compile(r"\b(mr|mrs|ms|miss|dr|prof)\.?\s+[A-Z][a-z]+", re.I)),
 ("name", re.compile(r"\b(name|patient|pt)\s*[:=]\s*[A-Z][a-z]+")),
 ("exact age >89", re.compile(r"\b(9\d|1[0-4]\d)\s*(y/?o|yrs?|years?[\s-]old)\b", re.I)),
]


class ImportError_(Exception):
    pass


def read_input(paths):
    """-> {sheet: [row dict]} with string/None cells."""
    data = {}
    if len(paths) == 1 and paths[0].lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(paths[0], read_only=True, data_only=True)
        for s in SHEETS:
            if s not in wb.sheetnames:
                continue
            it = wb[s].iter_rows(values_only=True)
            hdr = [str(h).strip() if h is not None else "" for h in next(it, [])]
            data[s] = [{h: ("" if v is None else str(v)).strip() for h, v in zip(hdr, r)} for r in it if any(v not in (None, "") for v in r)]
            data[s + "__header"] = hdr
        return data
    files = []
    for p in paths:
        files += [os.path.join(p, f) for f in sorted(os.listdir(p)) if f.endswith(".csv")] if os.path.isdir(p) else [p]
    for f in files:
        s = os.path.splitext(os.path.basename(f))[0]
        if s not in SHEETS:
            raise ImportError_(f"{f}: file name must be one of {', '.join(s2 + '.csv' for s2 in SHEETS)}")
        with open(f, newline="", encoding="utf-8-sig") as fh:
            rd = csv.DictReader(fh)
            data[s + "__header"] = [h.strip() for h in (rd.fieldnames or [])]
            data[s] = [{(k or "").strip(): (v or "").strip() for k, v in r.items()} for r in rd if any((v or "").strip() for v in r.values())]
    return data


def phi_hits(text):
    return [label for label, rx in PHI if rx.search(text)]


def validate(data):
    """Return list of error strings (empty = ok)."""
    errs = []
    if not any(s in data for s in CASE_SHEETS):
        return ["no Cases_Medical or Cases_NonMedical sheet found"]
    cases = {}
    for s, cols in SHEETS.items():
        if s not in data:
            continue
        hdr = data[s + "__header"]
        for h in hdr:
            if h and FORBIDDEN_COLS.search(h) and h not in cols:
                errs.append(f"{s}: column '{h}' looks like a direct identifier (PHI) and is not allowed")
            elif h and h not in cols:
                errs.append(f"{s}: unknown column '{h}' (allowed: {', '.join(cols)})")
        for c, (req, *_r) in cols.items():
            if req and c not in hdr:
                errs.append(f"{s}: missing required column '{c}'")
        for i, r in enumerate(data[s], 2):
            where = f"{s} row {i}"
            for c, v in r.items():
                if v and c not in ("options",):
                    for label in phi_hits(v):
                        errs.append(f"{where} [{c}]: possible PHI ({label}) in {v[:60]!r} — remove or generalise it")
            for c, (req, allowed, _d) in cols.items():
                v = r.get(c, "")
                if v == "":
                    if req:
                        errs.append(f"{where}: '{c}' is required")
                    continue
                if isinstance(allowed, list) and v not in allowed:
                    errs.append(f"{where}: '{c}'={v!r} not allowed (use one of: {', '.join(a for a in allowed if a)})")
                elif isinstance(allowed, str) and allowed.startswith("^") and not re.match(allowed, v):
                    errs.append(f"{where}: '{c}'={v!r} must match {allowed}")
                elif allowed in ("int", "int>=0"):
                    try:
                        n = int(float(v))
                        if allowed == "int>=0" and n < 0:
                            raise ValueError
                    except ValueError:
                        errs.append(f"{where}: '{c}'={v!r} must be {'a non-negative ' if allowed == 'int>=0' else 'an '}integer")
                elif allowed == "float":
                    try: float(v)
                    except ValueError: errs.append(f"{where}: '{c}'={v!r} must be a number")
                elif allowed == "prob":
                    try:
                        if not 0 <= float(v) <= 1: raise ValueError
                    except ValueError: errs.append(f"{where}: '{c}'={v!r} must be between 0 and 1")
                elif allowed == "json":
                    try:
                        if not isinstance(json.loads(v), dict): raise ValueError
                    except ValueError: errs.append(f"{where}: '{c}' must be a JSON object")
            if s in CASE_SHEETS and r.get("case_id"):
                if r["case_id"] in cases:
                    errs.append(f"{where}: duplicate case_id {r['case_id']}")
                cases[r["case_id"]] = r
                opts = OPTIONS_BY_DOMAIN.get(r.get("domain"), [])
                if r.get("options") and opts and r["options"].split(";") != opts:
                    errs.append(f"{where}: options must be exactly {';'.join(opts)} for {r.get('domain')}")
                if r.get("agent_proposal") and opts and r["agent_proposal"] not in opts:
                    errs.append(f"{where}: agent_proposal {r['agent_proposal']!r} is not one of the options")
    for s in ("Prior_Responses", "Outcomes", "Field_Reliability"):
        for i, r in enumerate(data.get(s, []), 2):
            c = cases.get(r.get("case_id"))
            if r.get("case_id") and c is None:
                errs.append(f"{s} row {i}: case_id {r.get('case_id')!r} not found in the case sheets"); continue
            if c is None:
                continue
            opts = OPTIONS_BY_DOMAIN.get(c.get("domain"), [])
            for col in ("human_answer", "outcome_value"):
                if r.get(col) and r[col] not in opts:
                    errs.append(f"{s} row {i}: {col} {r[col]!r} is not an option for {c['domain']}")
            if s == "Prior_Responses" and r.get("subject_key") and r["subject_key"] != c.get("subject_key"):
                errs.append(f"{s} row {i}: subject_key {r['subject_key']} does not match case {c['case_id']} ({c.get('subject_key')})")
    return errs


def num(v, f=int):
    try: return f(float(v)) if v not in ("", None) else None
    except ValueError: return None


def build(data):
    cases = {r["case_id"]: r for s in CASE_SHEETS for r in data.get(s, [])}
    outcomes = {r["case_id"]: r for r in data.get("Outcomes", [])}
    resp = {}
    for r in data.get("Prior_Responses", []):
        resp.setdefault(r["case_id"], []).append(r)
    fields = {}
    for r in data.get("Field_Reliability", []):
        fields.setdefault(r["case_id"], []).append(r)
    items = []
    for cid, c in cases.items():
        it = {"id": "IMP-" + cid, "case_id": cid, "domain": c["domain"], "context_id": c["context_id"], "subject_key": c["subject_key"],
              "day_offset": num(c["day_offset"]), "question": c["question"], "options": c["options"].split(";"),
              "agent_proposal": {"answer": c["agent_proposal"], "confidence": num(c["agent_confidence"], float)},
              "hard_trigger": c.get("hard_trigger") or None, "ambiguity": c["ambiguity"], "presentation": c["presentation"],
              "source": c["source"], "provenance": "imported",
              "fields": [{"field": f["field"], "reported": f.get("reported_value") or None, "truth": f.get("true_value") or None,
                          "state": f["reliability_state"], "app_state": RELIABILITY_APP[f["reliability_state"]], "info_source": f["info_source"]}
                         for f in fields.get(cid, [])]}
        if c["domain"] in ("ed_triage", "ai_tele_doc"):
            it.update(age_band=c["age_band"], sex=c["sex"], chief_complaint=c["chief_complaint"], channel=c["channel"],
                      vitals={k: num(c.get(k), float if k == "temp_c" else int) for k in ("hr", "sbp", "spo2", "temp_c")},
                      pmh=[p for p in (c.get("pmh") or "").split(";") if p])
        else:
            it.update(topic=c["topic"], extra=json.loads(c["extra_json"]) if c.get("extra_json") else {})
        items.append(it)
    return items, cases, outcomes, resp


def card_hash(payload):
    return hashlib.sha256(json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()


def decision_records(items, outcomes, resp):
    """Unified sdr.decision.v1 records for imported prior HITM responses. Unsigned by design (plan §4): they never enter the
    device's signed ledger; provenance=imported_unsigned keeps them separable in every export."""
    recs = []
    for it in items:
        o = outcomes.get(it["case_id"])
        payload = {k: it.get(k) for k in ("case_id", "subject_key", "presentation", "chief_complaint", "topic", "age_band", "sex",
                                          "vitals", "pmh", "fields", "extra", "ambiguity") if it.get(k) not in (None, [], {})}
        for r in resp.get(it["case_id"], []):
            recs.append({"schema": "sdr.decision.v1", "kind": "decision", "id": "import-" + r["response_id"], "timestamp": None,
                         "day_offset": it["day_offset"], "domain": it["domain"], "context_id": it["context_id"], "generation": 0,
                         "dial_snapshot": None, "card_payload": payload, "card_hash": card_hash(payload), "item_id": it["id"],
                         "question": it["question"], "options": it["options"], "agent_proposal": it["agent_proposal"],
                         "human_decision": {"answer": r["human_answer"], "label": r["human_decision"],
                                            "changed": r["human_answer"] != it["agent_proposal"]["answer"],
                                            "decided_by": "human (imported, role=" + r["reviewer_role"] + ")"},
                         "split_rating": r.get("split_rating") or None, "latency_ms": num(r.get("latency_ms")),
                         "reliability": {"label": "placeholder", "n": 0},
                         "escalation_flags": [it["hard_trigger"]] if it["hard_trigger"] else [],
                         "outcome": {"value": o["outcome_value"], "source": f"{o['outcome_source']} ({o['source']})", "observed_at": None,
                                     "observed_offset_days": num(o["observed_offset_days"]), "severity_if_wrong": o["severity_if_wrong"]} if o else None,
                         "signer_key_id": None, "signature": None, "ledger_hash": None, "ledger_prev": None,
                         "provenance": "imported_unsigned", "data_source": r["source"]})
    return recs


def encounters_csv(items, path):
    """AI tele-doc app importer format (EncounterImporter.columns). ED-triage complaints load there too (the AI doc review
    domain covers ED + outpatient presentations). Age band -> band midpoint (the app needs an Int; 90+ -> 90)."""
    from_app = ["encounter_id", "age", "sex", "chief_complaint", "channel", "companion", "interpreter", "language", "health_literacy",
                "hr", "sbp", "spo2", "temp", "pmh", "ai_plan", "transcript"] + [f"stated_{f}" for f in APP_FIELDS] + [f"true_{f}" for f in APP_FIELDS] + ["hard_trigger"]
    n = 0
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=from_app); w.writeheader()
        for it in items:
            if it["domain"] not in ("ed_triage", "ai_tele_doc"):
                continue
            lo, _, hi = it["age_band"].partition("-")
            age = int(lo.rstrip("+")) if not hi else (int(lo) + int(hi)) // 2
            v = it["vitals"]
            row = {"encounter_id": it["case_id"], "age": age, "sex": it["sex"], "chief_complaint": it["chief_complaint"], "channel": it["channel"],
                   "companion": "none", "interpreter": "no", "language": "English", "health_literacy": "",
                   "hr": v["hr"] or "", "sbp": v["sbp"] or "", "spo2": v["spo2"] or "", "temp": v["temp_c"] or "", "pmh": ";".join(it["pmh"]),
                   "ai_plan": f"{it['agent_proposal']['answer']} — {it['presentation']}", "transcript": it["presentation"],
                   "hard_trigger": it["hard_trigger"] or ""}
            for f in it["fields"]:
                if f["field"] in APP_FIELDS:
                    row[f"stated_{f['field']}"] = "—" if f["state"] == "not_said" else (f["reported"] or "")
                    row[f"true_{f['field']}"] = f["truth"] or ""
            w.writerow(row); n += 1
    return n


def dump_jsonl(path, rows):
    with open(path, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False, sort_keys=True) + "\n")
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def split(items, outcomes, resp, recs, out, cutoff):
    os.makedirs(out, exist_ok=True)
    train = [i for i in items if i["day_offset"] < cutoff]
    test = [i for i in items if i["day_offset"] >= cutoff]
    leaked = [r["response_id"] for i in test for r in resp.get(i["case_id"], [])]
    tid = {i["case_id"] for i in train}

    def hist(i):
        return {"case": i, "human_responses": [{k: r[k] for k in ("human_decision", "human_answer", "split_rating", "reviewer_role", "rationale_tag")}
                                                for r in resp.get(i["case_id"], [])],
                "outcome": ({k: outcomes[i["case_id"]][k] for k in ("outcome_value", "outcome_source", "observed_offset_days", "severity_if_wrong")}
                            if i["case_id"] in outcomes else None)}
    train_hist = [hist(i) for i in train]
    with_h, no_h, answers = [], [], []
    for i in sorted(test, key=lambda x: (x["day_offset"], x["case_id"])):
        same = [h for h in train_hist if h["case"]["subject_key"] == i["subject_key"] and h["case"]["day_offset"] < i["day_offset"]]
        pool = [h["case"]["case_id"] for h in train_hist if h["case"]["domain"] == i["domain"]]
        no_h.append({"condition": "no_history", "case": i})
        with_h.append({"condition": "with_history", "case": i, "subject_history": same, "domain_pool_case_ids": pool})
        o = outcomes.get(i["case_id"])
        answers.append({"case_id": i["case_id"], "outcome_value": o["outcome_value"] if o else None,
                        "outcome_source": o["outcome_source"] if o else None, "severity_if_wrong": o["severity_if_wrong"] if o else None})
    files = {
        "train/history.jsonl": train_hist,
        "train/decisions.jsonl": [r for r in recs if r["card_payload"]["case_id"] in tid],
        "test/cases.jsonl": [{"case": i} for i in sorted(test, key=lambda x: (x["day_offset"], x["case_id"]))],
        "test/answers.jsonl": answers,
        "conditions/no_history.jsonl": no_h,
        "conditions/with_history.jsonl": with_h,
    }
    sha = {}
    for rel, rows in files.items():
        p = os.path.join(out, rel); os.makedirs(os.path.dirname(p), exist_ok=True); sha[rel] = dump_jsonl(p, rows)
    m = {"schema": "sdr.case_experiment.v1", "created_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
         "cutoff_day_offset": cutoff, "split_rule": "train: day_offset < cutoff; test: day_offset >= cutoff (temporal, per-dataset)",
         "counts": {"train_cases": len(train), "test_cases": len(test), "train_decisions": len(files["train/decisions.jsonl"]),
                    "test_with_subject_history": sum(1 for w in with_h if w["subject_history"]),
                    "test_with_outcome": sum(1 for a in answers if a["outcome_value"])},
         "dropped_test_responses_to_prevent_leak": leaked,
         "conditions": {"no_history": "test case only", "with_history": "test case + same-subject prior cases/responses/outcomes + domain pool from train"},
         "blinded": "test/cases.jsonl and conditions/* carry no outcome; test/answers.jsonl is held out",
         "sha256": sha, "data_sources": sorted({i["source"] for i in items}),
         "note": "imported records are unsigned (provenance=imported_unsigned) and never enter the device ledger"}
    json.dump(m, open(os.path.join(out, "manifest.json"), "w"), indent=2)
    return m


def run(paths, out, do_split=False, cutoff=30, validate_only=False, export=False):
    data = read_input(paths)
    errs = validate(data)
    if errs:
        raise ImportError_("\n".join(errs))
    items, cases, outcomes, resp = build(data)
    if validate_only:
        return {"ok": True, "cases": len(items)}
    recs = decision_records(items, outcomes, resp)
    os.makedirs(os.path.join(out, "items"), exist_ok=True)
    by = {}
    for it in items:
        by.setdefault(it["domain"], []).append(it)
    for d, its in by.items():
        json.dump({"domain": d, "provenance": "imported", "items": its}, open(os.path.join(out, "items", f"{d}.json"), "w"), indent=1, ensure_ascii=False)
    dump_jsonl(os.path.join(out, "decisions_imported.jsonl"), recs)
    n_enc = encounters_csv(items, os.path.join(out, "encounters.csv"))
    res = {"ok": True, "cases": len(items), "by_domain": {d: len(v) for d, v in by.items()}, "decisions": len(recs),
           "outcomes": len(outcomes), "encounters_csv_rows": n_enc, "out": out}
    if export:
        import unified_export as ux
        res["export"] = ux.export(recs, os.path.join(out, "unified"), "imported_cases")
    if do_split:
        res["experiment"] = split(items, outcomes, resp, recs, out, cutoff)["counts"]
    return res


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("--out", default="build/case_import")
    ap.add_argument("--split", action="store_true")
    ap.add_argument("--cutoff", type=int, default=30)
    ap.add_argument("--validate-only", action="store_true")
    ap.add_argument("--export", action="store_true", help="also run unified_export (SFT/preference/RL rows/replay) on the imported records")
    a = ap.parse_args(argv)
    try:
        res = run(a.inputs, a.out, a.split, a.cutoff, a.validate_only, a.export)
    except ImportError_ as e:
        print("IMPORT REJECTED — nothing written:\n" + str(e), file=sys.stderr); return 2
    print(json.dumps(res, indent=2)); return 0


if __name__ == "__main__":
    sys.exit(main())
